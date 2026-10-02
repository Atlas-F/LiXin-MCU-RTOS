from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = '全球传统木制风帆船造船学习项目与申请联系清单.docx'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.1); sec.bottom_margin = Cm(1.9)
sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.25
for name, size in [('Title', 22), ('Heading 1', 15), ('Heading 2', 12)]:
    st = styles[name]
    st.font.name = 'Microsoft YaHei'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(10)
styles['Heading 1'].font.bold = True
styles['Heading 1'].paragraph_format.space_before = Pt(16)
styles['Heading 1'].paragraph_format.space_after = Pt(7)
styles['Heading 2'].font.bold = True
styles['Heading 2'].paragraph_format.space_before = Pt(10)
styles['Heading 2'].paragraph_format.space_after = Pt(4)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:'+side))
        if node is None:
            node = OxmlElement('w:'+side); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def set_cell_border(cell, color='D9D9D9', sz='8'):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ['top','left','bottom','right','insideH','insideV']:
        tag = qn('w:'+edge); elem = borders.find(tag)
        if elem is None:
            elem = OxmlElement('w:'+edge); borders.append(elem)
        elem.set(qn('w:val'),'single'); elem.set(qn('w:sz'),sz); elem.set(qn('w:color'),color)

def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell); set_cell_border(cell)
            if widths: cell.width = Cm(widths[ci])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.08
                for r in p.runs:
                    r.font.name = 'Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size = Pt(8.2)
                    if ri == 0: r.font.color.rgb = RGBColor(255,255,255); r.font.bold = True
            if ri == 0: set_cell_shading(cell,'163A5F')
            elif ri % 2 == 0: set_cell_shading(cell,'F4F7FA')

def add_table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for i,h in enumerate(headers): t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i,val in enumerate(row): cells[i].text = val
    style_table(t,widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def para(text='', style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); r.bold = True
        p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2); p.add_run(text); return p

def page_break(): doc.add_page_break()

# Cover
p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('全球传统木制风帆船造船学习项目与申请联系清单')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('面向传统木船 船匠技艺 历史帆船复原与海外学习的研究型行动指南'); r.italic=True; r.font.size=Pt(11)
doc.add_paragraph('')
para('这份文档把“可公开报名的学校”与“仍在运行但只能通过关系进入的船厂”明确分开。对想接近地理大发现时代欧洲木帆船的人，最重要的不是只学会造一条小艇，而是依次获得木工和放样基础、传统船体结构能力、桅杆与索具认知，再进入历史复原或大型木帆船船厂。')
para('核心结论：最具历史匹配度且可实际申请的是西班牙 Albaola；最清晰的英语国际学生签证路径是美国 Northwest School of Wooden Boatbuilding；英国 Boat Building Academy 适合从零建立基础。印尼 Pinisi、土耳其 Bodrum gulet、阿曼 Sur dhow 是重要的活态船厂，但目前没有证据表明它们对外国人提供可直接报名的规范学徒班。')
para('核验时间：2026年10月1日。费用、开课与签证政策会变化，提交申请或购票前必须向机构以书面方式二次确认。')

doc.add_heading('阅读方法', level=1)
para('A 类为“可直接申请”的项目；B 类为“可联系、但录取或跟工机会不能预先保证”的活态船厂/遗产项目；C 类适合短期试水或补充技能。本文中的“已核实”仅表示在公开官方资料或可靠项目资料中能看到当前项目、课程或联系方式，并不等于已为申请人保留名额。')

doc.add_heading('一 全球相关造船厂 项目与学校分析列表', level=1)
para('下表是按与目标的关系、实际可进入性和学习价值筛选的全球清单。并非“有传统船展示”的博物馆都会列入，只有能看到课程、真实建造/修复现场、学徒制或可实际联系的船厂才保留。')

rows = [
['A','西班牙 Pasaia','Albaola Aprendiztegi','16世纪巴斯克跨洋捕鲸船 San Juan 复原；传统船体、船木工、历史考证','3年制船匠学校；约18名学徒；面向全球18岁以上','最贴近大航海时代欧洲远洋木帆船'],
['A','美国 Port Hadlock','Northwest School of Wooden Boatbuilding','传统木船建造，兼顾现代船舶系统','12个月全日制职业课程；国际生 I-20 路径已公开','英语与签证最清晰'],
['A','英国 Lyme Regis','Boat Building Academy','放样、细木工、修复、clinker 与 carvel 铺板','40周全日制；另有1日至短期课程','零基础基础年'],
['A','美国 Newport','IYRS Boatbuilding and Restoration','古典木船建造、修复、表面处理','20个月全日制','系统修复能力强，成本高'],
['A','澳大利亚 Franklin','Wooden Boat Centre','传统 clinker/carvel 小船、放样、桨橹、木工','一年制及短课','英语实践强，签证资质待问'],
['A','法国 Douarnenez','Les Ateliers de l Enfer','船体木工、制帆、船用软装','职业培训与 CAP 船舶木工路线','法语是实质门槛'],
['A','挪威 Rissa','Fosen Folkehøgskole','从选树、锯材到北欧叠板船','全年制民众高等学校课程','传统材料与叠板法很强'],
['C','美国 Brooklin','WoodenBoat School','小型传统木船、放样、修复与帆船基础','多为1至2周密集课','低风险试水'],
['C','新西兰 Auckland','New Zealand Traditional Boat Building School','木船建造、修复、半模型、clinker 铺板','全年短期工坊；2026有公开排期','低成本试水与作品积累'],
['B','荷兰 Lelystad','Museum Batavialand','17世纪 VOC 商船 Batavia 复原、索具、雕刻、传统船木工','工作船厂与志愿者/工坊机会','历史船现场强，非正规学位'],
['B','希腊 Samos','Museum of Aegean Boatbuilding and Maritime Crafts 与本地船厂','爱琴海 kaiki、trechantiri','公开资料见2年学习+1年学徒设想；未确认正在国际招生','应联系确认，不应预付'],
['B','克罗地亚 Betina','Betina Museum of Wooden Shipbuilding','达尔马提亚小型木船传统','成人季节工坊与当地船匠网络','文化与短期实践价值高'],
['B','印尼 Bulukumba','Ara Bira Tanah Beru Pinisi 船厂群','大型 Pinisi 双桅木帆船；家族口传比例与船体优先建造','家族师徒/商业船厂；无可确认国际课程','大型活态木帆船价值极高'],
['B','土耳其 Bodrum','gulet 船厂群','大型木质 gulet，双桅/三桅旅游帆船','商业船厂；逐家询问观察或助理机会','大型木船和商业工艺'],
['B','阿曼 Sur','dhow 船厂','阿拉伯传统木质 dhow','私人船厂，适合考察与关系型接触','活态手工技艺，公开入口弱'],
]
add_table(['级别','地区','机构或项目','主要技艺与船型','参与形式','对目标的判断'],rows,[0.8,2.25,3.15,4.15,3.45,2.25])

doc.add_heading('1.1 最值得优先的正式学校', level=2)
para('Albaola 是本清单中历史匹配度最高的项目。它在西班牙巴斯克 Pasaia，以复原16世纪 Basque 跨洋捕鲸船 San Juan 为核心；该船与“地理大发现时代”处于同一海洋技术世界。公开资料显示，Aprendiztegi 面向来自世界各地、18岁以上的男女学徒，项目为三年制，约18名学徒；2024年公开招生信息称免学费。它的优势不是“做一条练习艇”，而是让学徒在真实的历史船复原体系里理解船体、材料、船木工与考证。限制是授课语言和非欧盟居留支持尚需逐项确认。')
para('Northwest School of Wooden Boatbuilding 是风险最低的英语职业教育选择。其公开招生页明确说明，国际学生可在获得 I-20 后支付 SEVIS 费并申请签证；2026项目已开放报名。课程为12个月全日制，适合作为进入木船行业的基础训练，而不是直接学习一艘16世纪大帆船的复原。')
para('Boat Building Academy 是英语零基础路径中的强选项：40周、每周五天全日制，公开资料可见放样、接合、建造与修复内容，并包含 clinker 与 carvel 等传统铺板方法。它最适合先把手工具、尺寸精度、放样与木船基本流程练成可用技能。')
para('IYRS 更偏“古典木船建造与修复”的完整职业训练。20个月周期和两年约 US$49,350 的学费决定了它适合预算充足、希望以后进入修复船厂的人。塔斯马尼亚 Wooden Boat Centre 同样是强实践选择，课程会涉及传统 clinker 和 carvel 小船，但需先确认是否能为国际学生提供可用的学生签证路径。')

doc.add_heading('1.2 活态船厂路线的价值与边界', level=2)
para('印尼 Pinisi、土耳其 gulet、阿曼 dhow 的价值在于它们不只是复原，而是仍在制造较大尺度木帆船。但它们的风险同样明确：你无法像申请大学一样在线提交表格、获得标准课程表或保证工位。进入这些现场通常依赖船期、船东预算、家族船匠关系、语言、保险和当地许可。应把它们当成第二阶段目标，而不是第一站。')
bullet('印尼 Bulukumba：Pinisi 被列入 UNESCO 非遗，核心区域为 Tanah Beru、Bira、Ara。知识多由 panrita lopi 船匠家族口传，常不依赖完整书面图纸。适合在已有木工/木船基础后，申请短期观察并逐步争取辅助工作。')
bullet('土耳其 Bodrum：gulet 是传统木质双桅船，Bodrum 仍有大型商业船厂。工艺中会混入现代游艇标准与材料，因此它更适合学“大型木船商业建造”，而不是纯历史复原。')
bullet('希腊 Samos：公开资料显示博物馆曾提出2年培训加1年学徒的构想，但未找到2026国际公开招生页。可以联系，不可以据此做留学承诺。')

doc.add_heading('二 分析建议与学习路径', level=1)
para('“会造木船”与“会造大航海时代帆船”之间至少隔着四层能力。不要跳过前两层直接去找大船船厂，否则大概率只能参观。')
add_table(['阶段','要获得的能力','可验证产出','优先地点'],[
['第一阶段 基础','安全使用手工具；木材识别；刨削、凿榫、量测；简单放样','完成小尺度半模型、框架或小艇部件；有清晰过程照片','NZTBS、WoodenBoat School、BBA短课'],
['第二阶段 船体','全尺寸放样；龙骨、肋骨、铺板；clinker/carvel；填缝与涂装','参与或独立完成一艘小型木船；能解释板缝、肋骨与木材纹理选择','BBA、Northwest、Wooden Boat Centre、Fosen'],
['第三阶段 修复与帆装','腐朽诊断、拆换结构、桅杆/圆材、索具与制帆协作','修复前后记录；能说明原构件保留与替换依据','IYRS、Les Ateliers、历史船博物馆'],
['第四阶段 历史大船','大尺度结构、史料考证、传统工法与团队协作','在真实复原/建造项目中承担可交付工序','Albaola、Batavialand；再联系 Pinisi/Bodrum'],
], [2.65,5.3,5.4,4.1])
para('推荐路线：先用短课验证兴趣，再选一年或更长的系统课程，最后进入历史船或活态大型船厂。若你愿意投入三年且语言能解决，Albaola 应排第一；若要稳妥英语留学，则 Northwest 排第一；若只先判断自己是否适合长期的体力劳动、木粉、重复刨削和尺寸精度，先上 NZTBS 或 WoodenBoat School 的短课。')
para('不建议把“去印尼或土耳其看一趟船厂”当作学习成果。最小可验证目标应是：你能独立解释一块船板为什么这样放样、为什么要蒸弯/冷弯、板缝如何防水、坏肋骨如何判断是否需更换，并留下自己实际参与的作品和照片。')

doc.add_heading('三 已落实的申请和联系信息清单', level=1)
para('以下是按当前可执行性排序的联系清单。每一项都写明了“已经确认什么”和“必须由对方书面确认什么”。在未获得录取或正式邀请前，不要购买不可退机票，也不要把学费、住宿或签证条件当作已确定。')
add_table(['优先级','机构与官方入口','已核实','待书面确认','申请人行动'],[
['1','Albaola Aprendiztegi\nalbaola.org/en/boatbuilding-school/','3年制；约18学徒；18岁以上；全球申请者；公开资料称免学费','下一届日期；授课语言；西语/巴斯克语要求；住宿；非欧盟居留支持','发英文信并附个人简介；同时问是否可安排线上沟通'],
['2','Northwest School of Wooden Boatbuilding\nnwswb.edu/programs/boatbuilding/','2026招生开放；12个月；国际学生 I-20 路径公开；2026-27每季度 US$6,950（学费+材料）','语言证明；住宿；I-20流程与资金证明；工具清单；最新总预算','在线申请前先向招生办确认国际申请材料和开学日'],
['3','Boat Building Academy\nboatbuildingacademy.com/40-week-boat-building-course/','40周；每周5天；公开资料显示课程费约 £17,000以上','对中国国籍可用的签证文件；最新学费；住宿；英语要求','先问签证可行性；若不支持长期签证则考虑短课'],
['4','IYRS\niyrs.edu/programs/boatbuilding-restoration/','20个月/76周；年学费 US$24,675，材料费另计；有国际学生群体','是否签发 I-20；F-1资格；住宿；英语证明；奖助','索要国际申请清单与全程成本'],
['5','Wooden Boat Centre\nwoodenboatcentre.com/','一年制与短课；公开未来一期价格 A$26,800','课程对应年份；是否 CRICOS 注册并可办学生签证；国际学员资格','不要按页面未来日期付款；先邮件确认近期开课'],
['6','Les Ateliers de l Enfer\nlesateliersdelenfer.fr/','自1984年专业培训；2026仍有招生宣讲；船体木工/制帆/软装','法语门槛；非欧盟入学与签证；学费；课程周期','若法语未到 B1/B2，先不要作为主申请'],
['7','Fosen Folkehøgskole\nfosen.fhs.no/boat-building/','全年传统造船；公开参考年费 NOK 150,310','下一期开课；挪威语/英语；国际学生签证；包含项目','发信问语言和居留；确认是否对非北欧成年申请者开放'],
['试水','New Zealand Traditional Boat Building School\nnztbs.org.nz/','2026有两天基础技能、半模型等工坊排期；公益信托模式','具体费用、工具、旅游签适用范围','用短课建立第一件作品，不用于长期就业签证'],
], [1.15,4.85,4.15,4.6,4.25])

doc.add_heading('3.1 船厂和遗产项目联系清单', level=2)
add_table(['地点','入口','可以合理请求的内容','不要直接假定的内容','建议的第一次联系'],[
['印尼 Bulukumba','bulukumbaboatbuilders.com/\nRiara Marine 电话 +62 811-3885-3333','开工船期、1-2周观察访问、是否可接收有基础者做辅助','正式课程、无偿长期岗位、旅游签下工作许可','说明你尊重当地工艺；问哪艘船何时在建、可否由向导陪同观察'],
['土耳其 Bodrum','Aganlar Tersanesi；Oguz Marin；当地 gulet 船厂','船厂参观、实习/助理机会、当前木船建造或修复项目','公开国际学徒班、纯传统历史工法','写明自己希望观察或辅助，不要一开始就要求“拜师几年”'],
['希腊 Samos','Museum of Aegean Boatbuilding and Maritime Crafts；Patmos Marine 等本地船厂线索','项目现状、工坊、志愿或短期观察','2年+1年计划已经对国际生正式招生','只问事实：下一期培训是否存在、何语授课、是否接受国际成人'],
['荷兰 Lelystad','Museum Batavialand\nbatavialand.nl/','志愿者机制、传统船厂日、工坊和活动日','有薪国际学徒或签证担保','询问是否有英语志愿/工坊以及工作船厂开放日'],
['阿曼 Sur','Sur dhow 相关船厂与当地文化/旅游机构','船厂参观、建造期、当地协调人','标准化培训或自由进入工作区','通过当地机构预约，尊重私人船厂与拍摄限制'],
], [2.1,4.2,4.5,4.6,4.15])

doc.add_heading('3.2 联系与申请的顺序', level=2)
para('第一周同时向 Albaola、Northwest、BBA 发信；这三家分别回答“历史匹配”“签证清晰”“基础训练”三个不同问题。邮件发送七天后可礼貌跟进一次。拿到回复后，用下表比较，不以名气或单纯学费做决定。')
add_table(['比较项','为什么重要','最优先问法'],[
['授课语言','车间安全与口头指导不能靠翻译软件','What language is used for daily workshop instruction and safety briefings?'],
['合法身份','课程能否转化为合法长期停留决定项目是否可行','Can you issue or support the documents needed for a student or training visa?'],
['真实工时','同为“一年”，手上干活的时间可能差很多','How many supervised workshop hours are included each week?'],
['课程对象','小艇、游艇修复与历史大船不是同一训练','Which traditional construction methods and vessel types are taught in the current intake?'],
['总成本','学费之外还有材料、工具、住宿、保险与签证','Please provide the estimated total cost including required tools and materials.'],
['成果归属','是否能留下作品集直接影响下一站申请','May students document their work and build a portfolio during the programme?'],
], [3.0,6.7,7.95])

doc.add_heading('四 对比与选择建议', level=1)
add_table(['如果你的首要条件是','优先选择','原因','不宜优先'],[
['最接近欧洲大航海时代历史船','Albaola','真实16世纪跨洋船复原；三年深度','只看短课'],
['英语和学生签证确定性','Northwest','公开 I-20 路径；2026招生开放','把关系型船厂作为第一站'],
['零基础但想先建立手艺','BBA 或 Wooden Boat Centre','有明确车间教学、基础方法、连续练习','直接去 Pinisi 或 Bodrum 求学徒'],
['古典木船修复职业方向','IYRS','20个月修复与建造训练','只选非正规参观项目'],
['预算极低且能解决语言与居留','Albaola','公开资料称免学费','假定住宿与签证也免费'],
['想亲眼见大型活态木帆船','印尼 Pinisi 或 Bodrum gulet','今天仍在造较大船体','把考察当成规范培训'],
], [3.3,3.8,7.2,3.35])
para('最终建议：把“学校”当作技能和合法身份的基础设施，把“活态船厂”当作第二阶段的深度现场。这样既能保留你对古船、风帆和世界各地船匠文化的兴趣，也不会把时间和签证风险押在没有公开招生机制的地方。')

doc.add_heading('五 英文申请或联系信', level=1)
para('使用方法：学校联系时保留第1-5项；船厂联系时保留第6项，并将 subject 改为 “Inquiry about a short observation or workshop assistant placement”。请把方括号内容替换为真实信息；不要夸大尚未具备的木工经验。')

email = [
('Subject: International applicant inquiry - traditional wooden boatbuilding', True),
('', False),
('Dear [Name or Admissions Team],', False),
('', False),
('My name is [Your Name], and I am a prospective applicant from China with a serious long-term interest in traditional wooden sailing vessels and historic shipbuilding.', False),
('', False),
('I am particularly interested in learning practical skills such as lofting, timber selection and shaping, framing, planking, caulking, restoration, spars, and traditional rigging. My longer-term goal is to contribute to the construction or restoration of historic wooden sailing ships.', False),
('', False),
('Could you please advise on the following?', False),
('1. Is your next [course / apprenticeship / placement] open to applicants from outside the EU, UK, or US?', False),
('2. What are the application deadline, duration, tuition, and total expected tool or material costs?', False),
('3. What language level is required for workshop instruction and safety?', False),
('4. Can your organisation provide documents or support needed for a student or training visa?', False),
('5. Is prior woodworking experience required? If so, what portfolio or evidence would you like to see?', False),
('6. For a shipyard: would a short observation visit or workshop-assistant placement be possible while a traditional wooden vessel is under construction?', False),
('', False),
('I would be grateful for any current prospectus, application instructions, and advice for an overseas applicant.', False),
('', False),
('Kind regards,', False),
('[Your Name]', False),
('[City, China]', False),
('[Email]', False),
]
t = doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.rows[0].cells[0].width=Cm(18.0); set_cell_shading(t.rows[0].cells[0],'F4F7FA'); set_cell_border(t.rows[0].cells[0],'B7C7D6'); set_cell_margins(t.rows[0].cells[0],180,180,180,180)
cell=t.rows[0].cells[0]; cell.text=''
for i,(text,bold) in enumerate(email):
    p = cell.paragraphs[0] if i==0 else cell.add_paragraph()
    p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.15
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(9.5)

doc.add_heading('六 核验基础与使用边界', level=1)
para('本报告基于 2026年10月1日前后公开可检索的官方页面、项目页面、博物馆/学校资料和必要的本地商家公开联系方式整理。费用与开课信息只作为申请前筛选用的参考值。尤其是免学费项目、未来年度课程、签证资格、船厂观察机会、住处和劳动许可，必须由机构用当前书面回复确认。')
para('主要核验入口：Albaola（albaola.org）；Northwest School of Wooden Boatbuilding（nwswb.edu）；Boat Building Academy（boatbuildingacademy.com）；IYRS（iyrs.edu）；Wooden Boat Centre（woodenboatcentre.com）；Les Ateliers de l Enfer（lesateliersdelenfer.fr）；Fosen Folkehøgskole（fosen.fhs.no）；NZ Traditional Boat Building School（nztbs.org.nz）；UNESCO Pinisi 条目（ich.unesco.org）；Museum Batavialand（batavialand.nl）。')

# footer page number field
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('全球传统木制风帆船造船学习项目与申请联系清单  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); run._r.addnext(fld)

doc.core_properties.title = '全球传统木制风帆船造船学习项目与申请联系清单'
doc.core_properties.subject = '传统木制风帆船造船学习项目研究与申请联系指南'
doc.core_properties.author = ''
doc.save(OUT)
print(OUT)
